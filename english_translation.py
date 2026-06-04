from openpyxl import load_workbook
from concurrent.futures import ThreadPoolExecutor, as_completed
from pypinyin import pinyin, Style
import re

# Load workbook
wb = load_workbook(r"D:\\JE project\\python\\special_char.xlsx")
ws = wb.active

# Headers
headers = [cell.value for cell in ws[1]]
name_idx = headers.index("name") + 1
desc_idx = headers.index("Description") + 1

# New columns
new_name_col = len(headers) + 1
new_desc_col = len(headers) + 2
ws.cell(1, new_name_col, "name_pinyin")
ws.cell(1, new_desc_col, "merge_desc_pinyin")

# Regex for Chinese characters
chinese_pattern = re.compile(r'[\u4e00-\u9fff]+')

# Cache for converted text
cache = {}

def chinese_to_iso7098(text):
    """
    Convert Chinese text to ISO 7098 (Hanyu Pinyin, no tones)
    """
    if not text:
        return text

    text = str(text)

    # Cache check
    if text in cache:
        return cache[text]

    # If no Chinese characters, return as-is
    if not chinese_pattern.search(text):
        cache[text] = text
        return text

    def convert(match):
        chars = match.group(0)
        result = pinyin(chars, style=Style.NORMAL, strict=False)
        return " ".join(item[0] for item in result)

    converted_text = chinese_pattern.sub(convert, text)

    cache[text] = converted_text
    return converted_text


def process_row(row):
    original_name = ws.cell(row, name_idx).value
    original_desc = ws.cell(row, desc_idx).value

    converted_name = chinese_to_iso7098(original_name)
    converted_desc = chinese_to_iso7098(original_desc)

    return row, original_name, converted_name, original_desc, converted_desc


# Multithreading
max_workers = 50
futures = []

with ThreadPoolExecutor(max_workers=max_workers) as executor:
    for row in range(2, ws.max_row + 1):
        futures.append(executor.submit(process_row, row))

    for future in as_completed(futures):
        row, original_name, converted_name, original_desc, converted_desc = future.result()

        # Write to Excel
        ws.cell(row, new_name_col, converted_name)
        ws.cell(row, new_desc_col, converted_desc)

        # Progress log
        print(f"Row {row} updated:")
        print(f"  Name: {original_name} -> {converted_name}")
        print(f"  Desc: {original_desc} -> {converted_desc}")
        print("-" * 50)

# Save output
wb.save(r"D:\\JE project\\python\\pinyin_iso7098.xlsx")
print("✅ Pinyin conversion (ISO 7098) completed!")